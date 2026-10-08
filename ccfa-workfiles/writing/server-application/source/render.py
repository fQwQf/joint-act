"""Render the editable server application from its canonical content and measured GPU evidence."""
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[4]
SOURCE = Path(__file__).with_name("report.json")
OUTPUT = ROOT / "docs/proposal/server-application.docx"
content = json.loads(SOURCE.read_text(encoding="utf-8"))
gpu = json.loads((ROOT / "docs/evidence/gpu3090-validation.json").read_text(encoding="utf-8"))
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.65), Cm(1.65)
section.left_margin, section.right_margin = Cm(2), Cm(2)
section.header_distance, section.footer_distance = Cm(.65), Cm(.65)
for name in ("Normal", "Heading 1", "Heading 2", "Title"):
    style = doc.styles[name]
    style.font.name = "Noto Sans CJK SC"
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
normal = doc.styles["Normal"]
normal.font.size = Pt(10)
normal.paragraph_format.line_spacing = 1.14
normal.paragraph_format.space_after = Pt(7)
for name in ("Heading 1", "Heading 2"):
    doc.styles[name].font.size = Pt(13)
    doc.styles[name].font.color.rgb = RGBColor.from_string("173C5E")
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(6)
header = section.header.paragraphs[0]
header.text = "JointAct  ·  研究开题与算力申请"
header.style = doc.styles["Normal"]
header.runs[0].font.size = Pt(8)
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run(content["date"] + "  ·  ").font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
plain = []
for index, blocks in enumerate(content["pages"]):
    if index:
        doc.add_page_break()
    else:
        title = doc.add_paragraph()
        run = title.add_run(content["title"])
        run.bold, run.font.size = True, Pt(21)
        run.font.color.rgb = RGBColor.from_string("173C5E")
        doc.add_paragraph(content["subtitle"])
    for kind, value in blocks:
        if kind == "gpu":
            value = (f"2026年10月8日，RTX 3090 24 GB已完成完整7B基座的{gpu['final_step']}步训练，"
                     f"包含从第2步恢复的{gpu['resumed_steps']}步；双相机224像素、H=8、K=64、LoRA rank 16、"
                     f"单卡batch=1、梯度累积16、BF16。记录峰值分配显存{gpu['peak_allocated_gib']:.2f} GiB，"
                     f"峰值预留显存{gpu['peak_reserved_gib']:.2f} GiB；恢复段平均优化耗时"
                     f"{gpu['resumed_seconds_per_step']:.2f}秒/步。导出后对16个保留观测进行32次推理测量，"
                     f"每次8步动作预测的端到端中位延迟为{gpu['benchmark']['median_ms']:.1f}毫秒。"
                     "短训及保存完成，完整收敛和长时间运行尚需检验。")
        if kind == "budget":
            value = (f"3090恢复段短测为{gpu['resumed_seconds_per_step']:.2f}秒/优化步。按相同配置线性外推，"
                     f"20,000步约需{gpu['estimated_20000_step_hours']:.1f} GPU小时，12组内部对照约需"
                     f"{12 * gpu['estimated_20000_step_hours']:.0f} GPU小时。此估算不含启动、定期验证、"
                     "保存和闭环评测；短测未覆盖完整训练周期。更高端GPU的加速比例需在实际服务器上测定。")
        if kind == "table":
            table = doc.add_table(rows=0, cols=len(value[0]))
            table.style = "Light Shading Accent 1"
            table.autofit = False
            table.columns[0].width, table.columns[1].width = Cm(4), Cm(13)
            for ri, row in enumerate(value):
                cells = table.add_row().cells
                cells[0].width, cells[1].width = Cm(4), Cm(13)
                for cell, text in zip(cells, row):
                    cell.text = text
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after = Pt(4)
                        para.paragraph_format.space_before = Pt(4)
                        for run in para.runs:
                            run.font.size = Pt(9)
                            run.bold = ri == 0
                trpr = table.rows[-1]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:cantSplit"))
                plain.append(" | ".join(row))
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        else:
            p = doc.add_paragraph(value, "Heading 1" if kind == "h" else "Normal")
            if kind == "ref":
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(8)
            plain.append(value)
doc.core_properties.title = content["title"].replace("\n", "")
doc.core_properties.subject = content["subtitle"]
doc.core_properties.author = ""
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
build = SOURCE.parent.parent / "build"
build.mkdir(parents=True, exist_ok=True)
(build / "report.txt").write_text("\n\n".join(plain), encoding="utf-8")
print(OUTPUT)
